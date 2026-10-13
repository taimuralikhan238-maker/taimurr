"""ATS Resume Analyzer - Streamlit UI + Gemini Flash.

Upload a resume (PDF / DOCX / TXT), optionally paste a job description,
and get an estimated ATS score plus concrete improvement suggestions.
"""

import io
import json
import os
import re
from typing import Any, Dict, List, Optional

import streamlit as st
from docx import Document
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash")
MAX_FILE_MB = 5
MAX_CHARS = 20000  # resume/JD text sent to the model is truncated to this size
MIN_CHARS = 150  # below this we assume the file has no readable text

BREAKDOWN_LABELS = {
    "keywords": "Keywords & relevance",
    "formatting": "ATS-friendly formatting",
    "structure": "Sections & structure",
    "content_quality": "Content & impact",
    "readability": "Readability & grammar",
}

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume coach. You evaluate resumes the way modern ATS parsers and \
recruiters do. Be honest, specific and constructive. Never invent facts that are \
not in the resume. Respond with valid JSON only."""


# ----------------------------------------------------------------------------
# Text extraction
# ----------------------------------------------------------------------------
def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # many resumes put content in tables
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ----------------------------------------------------------------------------
# Prompt + response handling
# ----------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd = job_description.strip()
    if jd:
        jd_block = (
            "A target job description is provided. Score keyword match and relevance "
            "against it, and list important keywords from it that the resume is missing.\n\n"
            f"JOB DESCRIPTION:\n\"\"\"\n{jd[:MAX_CHARS]}\n\"\"\"\n"
        )
    else:
        jd_block = (
            "No job description is provided. Evaluate the resume for general ATS "
            "readiness and infer the most likely target role from its content.\n"
        )

    return f"""Analyze the resume below.

{jd_block}
RESUME:
\"\"\"
{resume_text[:MAX_CHARS]}
\"\"\"

Return ONLY a JSON object with exactly this structure (no markdown, no commentary):
{{
  "overall_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "breakdown": {{
    "keywords": <integer 0-100>,
    "formatting": <integer 0-100>,
    "structure": <integer 0-100>,
    "content_quality": <integer 0-100>,
    "readability": <integer 0-100>
  }},
  "strengths": ["<short point>", "..."],
  "weaknesses": ["<short point>", "..."],
  "missing_keywords": ["<keyword or skill>", "..."],
  "improvements": [
    {{"section": "<resume section>", "issue": "<what is wrong>", "suggestion": "<specific fix>"}}
  ],
  "rewritten_bullets": [
    {{"original": "<weak bullet taken from the resume>", "improved": "<stronger version>"}}
  ]
}}

Scoring guidance:
- Be realistic: an average resume scores 55-70; 85+ is rare.
- overall_score should reflect the breakdown scores.
- Give 4-8 improvements ordered by impact, and 2-4 rewritten bullets.
- Rewrites must use only facts present in the resume. If a metric is missing, use a
  placeholder like [X%] instead of making up numbers.
"""


def parse_json_response(text: str) -> Dict[str, Any]:
    """Parse JSON from a model response, tolerating code fences / stray text."""
    if not text or not text.strip():
        raise ValueError("The model returned an empty response.")
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start, end = cleaned.find("{"), cleaned.rfind("}")
        if start != -1 and end > start:
            return json.loads(cleaned[start : end + 1])
        raise


def _clamp_score(value: Any) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def _str_list(value: Any) -> List[str]:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalize_result(raw: Dict[str, Any]) -> Dict[str, Any]:
    """Coerce the model output into a predictable shape so the UI never crashes."""
    if not isinstance(raw, dict):
        raise ValueError("Unexpected response format from the model.")

    breakdown_raw = raw.get("breakdown") if isinstance(raw.get("breakdown"), dict) else {}
    breakdown = {key: _clamp_score(breakdown_raw.get(key)) for key in BREAKDOWN_LABELS}

    overall = raw.get("overall_score")
    if overall is None:
        scores = [v for v in breakdown.values() if v]
        overall = sum(scores) / len(scores) if scores else 0

    improvements = []
    for item in raw.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "section": str(item.get("section", "General")).strip() or "General",
                    "issue": str(item.get("issue", "")).strip(),
                    "suggestion": str(item.get("suggestion", "")).strip(),
                }
            )

    bullets = []
    for item in raw.get("rewritten_bullets") or []:
        if isinstance(item, dict) and item.get("improved"):
            bullets.append(
                {
                    "original": str(item.get("original", "")).strip(),
                    "improved": str(item.get("improved", "")).strip(),
                }
            )

    return {
        "overall_score": _clamp_score(overall),
        "summary": str(raw.get("summary", "")).strip(),
        "breakdown": breakdown,
        "strengths": _str_list(raw.get("strengths")),
        "weaknesses": _str_list(raw.get("weaknesses")),
        "missing_keywords": _str_list(raw.get("missing_keywords")),
        "improvements": improvements,
        "rewritten_bullets": bullets,
    }


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


# ----------------------------------------------------------------------------
# Gemini call
# ----------------------------------------------------------------------------
def analyze_resume(
    api_key: str, model: str, resume_text: str, job_description: str = ""
) -> Dict[str, Any]:
    """Send the resume to Gemini and return a normalized result dict."""
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        temperature=0.2,
        response_mime_type="application/json",
    )
    prompt = build_prompt(resume_text, job_description)

    last_error: Optional[Exception] = None
    for _ in range(2):  # one retry if the model returns malformed JSON
        response = client.models.generate_content(model=model, contents=prompt, config=config)
        try:
            return normalize_result(parse_json_response(response.text))
        except (ValueError, json.JSONDecodeError) as exc:
            last_error = exc
    raise ValueError(f"Could not read the model's response. Please try again. ({last_error})")


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def get_api_key() -> str:
    """Look for the key in Streamlit secrets, then the environment."""
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:  # no secrets file present
        key = ""
    return key or os.getenv("GEMINI_API_KEY", "")


def render_results(result: Dict[str, Any]) -> None:
    score = result["overall_score"]

    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("Estimated ATS score", f"{score}/100", score_label(score))
    with col2:
        st.progress(score / 100)
        if result["summary"]:
            st.write(result["summary"])

    st.subheader("Score breakdown")
    for key, label in BREAKDOWN_LABELS.items():
        value = result["breakdown"][key]
        st.write(f"**{label}** - {value}/100")
        st.progress(value / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for item in result["strengths"] or ["No strengths listed."]:
            st.markdown(f"- {item}")
    with right:
        st.subheader("Weaknesses")
        for item in result["weaknesses"] or ["No weaknesses listed."]:
            st.markdown(f"- {item}")

    if result["missing_keywords"]:
        st.subheader("Missing keywords")
        st.write(", ".join(f"`{kw}`" for kw in result["missing_keywords"]))

    if result["improvements"]:
        st.subheader("Suggested improvements")
        for i, item in enumerate(result["improvements"], start=1):
            with st.expander(f"{i}. {item['section']}", expanded=i <= 3):
                if item["issue"]:
                    st.markdown(f"**Issue:** {item['issue']}")
                if item["suggestion"]:
                    st.markdown(f"**Fix:** {item['suggestion']}")

    if result["rewritten_bullets"]:
        st.subheader("Example bullet rewrites")
        for item in result["rewritten_bullets"]:
            if item["original"]:
                st.markdown(f"**Before:** {item['original']}")
            st.markdown(f"**After:** {item['improved']}")
            st.divider()

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="ATS Resume Analyzer", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Analyzer")
    st.caption(
        "Upload your resume to get an estimated ATS score and tips to improve it. "
        "The score is an AI-generated estimate - real ATS software varies by employer."
    )

    api_key = get_api_key()
    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input(
                "Gemini API key",
                type="password",
                help="Get a free key at https://aistudio.google.com/apikey",
            )
        else:
            st.success("API key loaded")
        model = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.markdown("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Target job description (optional, but gives a more accurate score)",
        height=160,
        placeholder="Paste the job posting here...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            return
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading your resume..."):
                resume_text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(resume_text) < MIN_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned image, "
                "ATS systems cannot read it either - export it as a text-based PDF or DOCX."
            )
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(api_key, model.strip() or DEFAULT_MODEL, resume_text, job_description)
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            return

        st.session_state["result"] = result

    if "result" in st.session_state:
        st.divider()
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
